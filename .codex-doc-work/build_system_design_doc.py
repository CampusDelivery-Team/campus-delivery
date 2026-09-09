from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"D:\delivery-backend")
REFERENCE = ROOT / ".codex-doc-work" / "系统设计与实现文档模板.docx"
CONTENT = ROOT / ".codex-doc-work" / "system_design_content.md"
OUTPUT = ROOT / "docs" / "校园中转分发与跑腿服务管理系统-系统设计与实现文档.docx"


def font(run, east="宋体", latin="Times New Roman", size=10.5, bold=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for key, value in (("eastAsia", east), ("ascii", latin), ("hAnsi", latin)):
        rf.set(qn(f"w:{key}"), value)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    el = tcpr.find(qn("w:shd"))
    if el is None:
        el = OxmlElement("w:shd")
        tcpr.append(el)
    el.set(qn("w:fill"), fill)


def margins(cell, top=80, start=100, bottom=80, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = mar.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def cell_border(cell, edge, color="666666", size="5"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    el = borders.find(qn(f"w:{edge}"))
    if el is None:
        el = OxmlElement(f"w:{edge}")
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:color"), color)


def table_borders(table):
    pr = table._tbl.tblPr
    borders = pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), "808080")
        borders.append(el)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table_borders(table)
    header = table.rows[0]
    trpr = header._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    rep.set(qn("w:val"), "true")
    trpr.append(rep)
    for i, value in enumerate(rows[0]):
        cell = header.cells[i]
        shade(cell, "D9E2F3")
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        font(r, east="黑体", latin="Arial", size=9.5, bold=True)
    for ri, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            margins(cell)
            if ri % 2:
                shade(cell, "F7F9FC")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(value)
            font(r, size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(doc, text):
    p = doc.add_paragraph(style="文档正文" if "文档正文" in [s.name for s in doc.styles] else "Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text))


def add_heading(doc, text, level):
    if text.startswith("附录"):
        appendix_title = re.sub(r"^附录\s*A\s*", "", text).strip()
        p = doc.add_paragraph(appendix_title, style="附录" if "附录" in [s.name for s in doc.styles] else "Heading 1")
        p.paragraph_format.keep_with_next = True
        ppr = p._p.get_or_add_pPr()
        outline = ppr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        outline.set(qn("w:val"), "0")
        for run in p.runs:
            font(run, east="黑体", latin="Arial", size=15, bold=True)
        return
    if text in ("图索引", "表索引"):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        font(p.add_run(text), east="黑体", latin="Arial", size=14, bold=True)
        return
    text = re.sub(r"^\d+(?:\.\d+)*\s+", "", text)
    if level == 1 and len([p for p in doc.paragraphs if p.style.name == "Heading 1"]) > 0:
        doc.add_page_break()
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    for run in p.runs:
        font(run, east="黑体", latin="Arial", size={1: 15, 2: 14, 3: 12}[level], bold=True)


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(21)
    p.paragraph_format.first_line_indent = Pt(-10.5)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run("•  " + text))


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.right_indent = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    ppr.append(shd)
    font(p.add_run(text), east="等线", latin="Consolas", size=8.5)


def add_picture(doc, rel_path, caption, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ROOT / rel_path), width=Inches(float(width)))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(2)
    c.paragraph_format.space_after = Pt(5)
    font(c.add_run(caption), east="黑体", latin="Arial", size=10)


def clear(container):
    for child in list(container._element):
        container._element.remove(child)


def section_setup(section, page_format="decimal", start=None, show_page=True):
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.59)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.79)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    clear(section.header)
    ht = section.header.add_table(rows=1, cols=2, width=Inches(6.45))
    ht.autofit = False
    for i, text in enumerate(("校园跑腿系统", "设计与实现文档")):
        cell = ht.cell(0, i)
        margins(cell, 0, 0, 45, 0)
        cell_border(cell, "bottom")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        font(p.add_run(text), size=9.5)
    section.header.add_paragraph()
    clear(section.footer)
    ft = section.footer.add_table(rows=1, cols=2, width=Inches(6.45))
    ft.autofit = False
    for cell in ft.rows[0].cells:
        margins(cell, 45, 0, 0, 0)
        cell_border(cell, "top")
    lp = ft.cell(0, 0).paragraphs[0]
    font(lp.add_run("同济大学软件学院\n数据库课程设计项目"), size=8.5)
    rp = ft.cell(0, 1).paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if show_page:
        begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
        val = OxmlElement("w:t"); val.text = "1"
        end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
        run = rp.add_run()
        run._r.extend([begin, instr, sep, val, end])
        font(run, size=8.5)
    section.footer.add_paragraph()
    if start is not None:
        pg = section._sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            section._sectPr.append(pg)
        pg.set(qn("w:fmt"), page_format)
        pg.set(qn("w:start"), str(start))


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    font(p.add_run("目  录"), east="黑体", size=20, bold=True)
    p = doc.add_paragraph()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    val = OxmlElement("w:t"); val.text = "目录将在打开文档时自动更新"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run = p.add_run()
    run._r.extend([begin, instr, sep, val, end])
    font(run)


def parse_content(doc):
    lines = CONTENT.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith("[[IMAGE|"):
            parts = line[8:-2].split("|")
            add_picture(doc, parts[0], parts[1], parts[2])
        elif line.startswith("### "):
            add_heading(doc, line[4:], 3)
        elif line.startswith("## "):
            add_heading(doc, line[3:], 2)
        elif line.startswith("# "):
            add_heading(doc, line[2:], 1)
        elif line.startswith("- "):
            add_bullet(doc, line[2:])
        elif line.startswith("> "):
            add_code(doc, line[2:])
        else:
            add_body(doc, line)
        i += 1


def build():
    shutil.copyfile(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for name, size in (("Heading 1", 15), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")

    cover = doc.sections[0]
    section_setup(cover, show_page=False)
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("校园中转分发与跑腿服务管理系统"), east="黑体", size=24, bold=True)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("系统设计与实现文档"), east="黑体", size=20, bold=True)
    p.paragraph_format.space_after = Pt(34)
    members = [
        "2452207  刘相成", "2456179  唐独彪", "2451200  桑治", "2452326  崔少坤", "2453619  薛毓哲",
        "2452098  赵崇治", "2451347  韩昊苏", "2452281  李柏言", "2450299  谢智行", "2450333  蒋昊沄",
    ]
    for value in members:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        font(p.add_run(value), east="黑体", size=12, bold=True)
    toc = doc.add_section(WD_SECTION.NEW_PAGE)
    section_setup(toc, "upperRoman", 1, True)
    add_toc(doc)
    main = doc.add_section(WD_SECTION.NEW_PAGE)
    section_setup(main, "decimal", 1, True)
    parse_content(doc)
    props = doc.core_properties
    props.title = "校园中转分发与跑腿服务管理系统 系统设计与实现文档"
    props.subject = "数据库课程设计"
    props.author = "刘相成、唐独彪、桑治、崔少坤、薛毓哲、赵崇治、韩昊苏、李柏言、谢智行、蒋昊沄"
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
